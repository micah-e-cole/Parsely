# search.py
import re
from pathlib import Path
import openpyxl
from pdf2image import convert_from_path
import pytesseract
import pdfplumber           # preferred over PyPDF2
import docx                 # from python-docx


def run(pattern, target, ignore_case=True, extension="", pdf_mode="text"):
    """
    Generator that yields lines or cells matching pattern in a file or directory.
    Compatible with both CLI and GUI frontends.
    Supports plain text files, PDF (.pdf) files, and Excel (.xlsx) spreadsheets.
    """
    flags = re.IGNORECASE if ignore_case else 0
    regex = re.compile(pattern, flags)
    path = Path(target)

    if not path.exists():
        yield f"Error: {target} does not exist."
        return

    if extension and not extension.startswith("."):
        extension = "." + extension

    # If target is a single file
    if path.is_file():
        yield from _search_file(path, regex, pdf_mode)

    # If target is a directory
    elif path.is_dir():
        for file in path.rglob(f"*{extension}" if extension else "*"):
            if file.is_file():
                yield from _search_file(file, regex, pdf_mode)
    else:
        yield f"Error: {target} is not a file or directory."



def _search_file(file_path, regex, pdf_mode="text"):
    """
    Helper that yields matches from a single file.
    Detects .xlsx files and uses openpyxl to inspect cell values.
    Issues: Tesserect needs to be installed on Windows - no support for this right now
    """
    suffix = file_path.suffix.lower()

    if suffix == ".xlsx":
        yield from _search_excel(file_path, regex)
    # elif suffix == ".pdf":
    #     if pdf_mode == "scanned":
    #        yield from _search_pdf_ocr(file_path, regex)
    #     else:
    #        yield from _search_pdf(file_path, regex)
    elif suffix == ".pdf":
        yield from _search_pdf(file_path, regex)
    elif suffix == ".docx":
        yield from _search_docx(file_path, regex)
    else:
        yield from _search_text(file_path, regex)


# -------------------------
# Text file search
# -------------------------
def _search_text(file_path, regex):
    """Yield matching lines from a text-based file."""
    try:
        with open(file_path, encoding="utf-8", errors="ignore") as f:
            for i, line in enumerate(f, start=1):
                if regex.search(line):
                    yield f"{i} {line.strip()}"
    except Exception as e:
        yield f"Error reading {file_path}: {e}"


# -------------------------
# Excel file search
# -------------------------
def _search_excel(file_path, regex):
    """Yield matching cells from an Excel (.xlsx) workbook."""
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                for c_idx, cell in enumerate(row, start=1):
                    if cell and isinstance(cell, str) and regex.search(cell):
                        cell_ref = f"{openpyxl.utils.get_column_letter(c_idx)}{r_idx}"
                        yield f"Sheet: {sheet} | Cell {cell_ref} | {cell}"
        wb.close()
    except Exception as e:
        yield f"Error reading {file_path}: {e}"


# -------------------------
# PDF file search: Optimal, uses pdfplumber
# -------------------------
def _search_pdf(file_path, regex):
    """Accurate PDF search using pdfplumber."""
    try:
        with pdfplumber.open(file_path) as pdf:
            for page_num, page in enumerate(pdf.pages, start=1):
                text = page.extract_text() or ""
                text = re.sub(r"\s+", " ", text)
                for match in regex.finditer(text):
                    start = max(0, match.start() - 40)
                    end = match.end() + 40
                    snippet = text[start:end]
                    yield f"Page {page_num}: {snippet.strip()}"
    except Exception as e:
        yield f"Error reading {file_path}: {e}"


# -------------------------
# PDF file search, For scanned documents, uses OCR
# -------------------------
def _search_pdf_ocr(file_path, regex):
    """Perform OCR on each PDF page and search text."""
    try:
        images = convert_from_path(file_path)
        for page_num, img in enumerate(images, start=1):
            text = pytesseract.image_to_string(img)
            text = re.sub(r"\s+", " ", text)
            for match in regex.finditer(text):
                start = max(0, match.start() - 40)
                end = match.end() + 40
                snippet = text[start:end]
                yield f"Page {page_num} (OCR): {snippet.strip()}"
    except Exception as e:
        yield f"Error processing OCR for {file_path}: {e}"


# -------------------------
# Word file search
# -------------------------
def _search_docx(file_path, regex):
    """Yield matching paragraphs or table cells from Word (.docx) files."""
    try:
        document = docx.Document(file_path)
        # Search paragraphs
        for i, para in enumerate(document.paragraphs, start=1):
            text = para.text.strip()
            if text and regex.search(text):
                yield f"Paragraph {i}: {text}"

        # Search inside tables
        for t_idx, table in enumerate(document.tables, start=1):
            for r_idx, row in enumerate(table.rows, start=1):
                for c_idx, cell in enumerate(row.cells, start=1):
                    text = cell.text.strip()
                    if text and regex.search(text):
                        yield f"Table {t_idx}, Row {r_idx}, Col {c_idx}: {text}"
    except Exception as e:
        yield f"Error reading {file_path}: {e}"